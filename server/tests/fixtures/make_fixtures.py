"""Builds the server test fixtures (committed). Needs: python-docx, openpyxl, python-pptx, PyMuPDF.

    python server/tests/fixtures/make_fixtures.py
"""

import io
import zipfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
AR = "تقرير الصيانة الشهري"
BODY = "تم فحص جميع الأجهزة في المبنى الرئيسي ولا توجد أعطال."


def word():
    d = Document()
    d.add_heading(AR, 1)
    d.add_paragraph(BODY)
    t = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): "الجهاز", (0, 1): "الحالة", (1, 0): "مكيف", (1, 1): "سليم"}.items():
        t.cell(r, c).text = v
    d.save(HERE / "report.docx")


def excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "المبيعات"
    ws.sheet_view.rightToLeft = True
    ws.append(["الشهر", "المبلغ"])
    for m, v in [("يناير", 1500), ("فبراير", 2750), ("مارس", 3100)]:
        ws.append([m, v])
    ws.append(["الإجمالي", "=SUM(B2:B4)"])
    wb.save(HERE / "sales.xlsx")


def slides():
    p = Presentation()
    for title in (AR, "Office Lab"):
        s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = title
        box = s.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
        box.text_frame.text = BODY
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    p.save(HERE / "slides.pptx")


def pdfs():
    # Two text pages (Arabic needs a real font: Noto Naskh from the web app's node_modules is not here, so
    # use the PDF base font for Latin and check Arabic through the Office conversions instead).
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 100), f"Office Lab test page {i + 1}", fontsize=20)
        page.insert_text((72, 140), "The quick brown fox jumps over the lazy dog.", fontsize=12)
    doc.save(HERE / "text.pdf")

    # A heavy "photo" page: a smooth gradient stored at high resolution, so Ghostscript can shrink it.
    size = (1500, 2000)
    photo = Image.merge("RGB", [
        Image.blend(Image.linear_gradient("L").resize(size), Image.effect_noise(size, 30), 0.2),
        Image.blend(Image.radial_gradient("L").resize(size), Image.effect_noise(size, 20), 0.15),
        Image.blend(Image.linear_gradient("L").rotate(90).resize(size), Image.effect_noise(size, 25), 0.2),
    ])  # fmt: skip
    jpg = io.BytesIO()
    photo.save(jpg, "JPEG", quality=95)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=jpg.getvalue())
    doc.save(HERE / "photo.pdf")

    # Locked with an open password.
    doc = fitz.open(HERE / "text.pdf")
    doc.save(HERE / "locked.pdf", encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")

    # Damaged: wrong xref offsets and the tail cut off.
    raw = (HERE / "text.pdf").read_bytes()
    cut = raw[: raw.rfind(b"startxref")] + b"startxref\n99999\n%%EOF\n"
    (HERE / "damaged.pdf").write_bytes(cut.replace(b"0000000", b"0000999", 3))


def text_files():
    (HERE / "page.html").write_text(
        f'<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>{AR}</title>'
        f'<style>body{{font-family:"Noto Naskh Arabic",serif;font-size:20px}}</style></head>'
        f"<body><h1>{AR}</h1><p>{BODY}</p></body></html>",
        encoding="utf-8",
    )
    (HERE / "notes.md").write_text(f"# {AR}\n\n{BODY}\n\n- بند أول\n- بند ثان\n\n| الجهاز | الحالة |\n|---|---|\n| مكيف | سليم |\n",
                                   encoding="utf-8")
    (HERE / "paper.tex").write_text(
        "\\documentclass{article}\n\\title{Office Lab}\n\\begin{document}\n\\maketitle\n\\section{Intro}\n"
        "Energy is $E=mc^2$.\n\\begin{itemize}\\item one\\item two\\end{itemize}\n\\end{document}\n",
        encoding="utf-8",
    )


def hostile():
    """Inputs that try to read server files or reach internal services: must not leak anything."""
    (HERE / "hostile.tex").write_text(
        "\\documentclass{article}\\begin{document}\nA\\input{/etc/passwd}B\\include{/etc/hostname}\n\\end{document}\n",
        encoding="utf-8",
    )
    (HERE / "hostile.html").write_text(
        "<!doctype html><html><body><h1>probe</h1>"
        '<img src="http://api:8000/api/capabilities" onerror="document.body.dataset.e1=1">'
        '<iframe src="file:///etc/passwd" width="600" height="300"></iframe>'
        '<iframe src="http://api:8000/api/capabilities" width="600" height="300"></iframe>'
        '<iframe src="http://169.254.169.254/latest/meta-data/" width="600" height="300"></iframe>'
        "</body></html>",
        encoding="utf-8",
    )
    (HERE / "hostile.md").write_text("# probe\n\n![x](/etc/passwd)\n\n![y](file:///etc/hostname)\n", encoding="utf-8")
    # An EPUB whose chapter points at server files.
    chapter = (
        '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>c</title>'
        '<link rel="stylesheet" href="file:///etc/passwd"/></head><body><h1>probe</h1>'
        '<p>before</p><img src="file:///etc/passwd"/><img src="/etc/hostname"/><img src="../../../../etc/passwd"/>'
        '<object data="/etc/passwd"></object><p>after</p></body></html>'
    )
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">x</dc:identifier>'
        "<dc:title>probe</dc:title><dc:language>en</dc:language></metadata>"
        '<manifest><item id="c" href="c.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest>'
        '<spine><itemref idref="c"/></spine></package>'
    )
    nav = ('<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
           '<head><title>n</title></head><body><nav epub:type="toc"><ol><li><a href="c.xhtml">c</a></li></ol></nav>'
           "</body></html>")
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 "</rootfiles></container>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip")
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/c.xhtml", chapter)
        z.writestr("OEBPS/nav.xhtml", nav)
    (HERE / "hostile.epub").write_bytes(buf.getvalue())


if __name__ == "__main__":
    word()
    excel()
    slides()
    pdfs()
    text_files()
    hostile()
    for p in sorted(HERE.iterdir()):
        if p.suffix != ".py":
            print(f"{p.stat().st_size:>9}  {p.name}")
