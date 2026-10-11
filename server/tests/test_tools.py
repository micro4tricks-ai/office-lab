"""Every server tool against real fixtures. Runs inside the API container next to Gotenberg:

    docker build --target test -t office-lab-api:test server
    docker run --rm --network office-lab_internal office-lab-api:test        (stack up first)
"""

import io
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

fitz = pytest.importorskip("pymupdf")  # in the container (pdf2docx brings it)

from app import main
from app.config import settings

pytestmark = pytest.mark.tools
FIX = Path(__file__).parent / "fixtures"
SECRETS = ("root:x:0:0", "/usr/sbin/nologin", '"tools"', "ami-id", "maxMB")

@pytest.fixture(scope="module")
def client():
    if not shutil.which("pandoc"):
        pytest.skip("converters not installed (run inside the API container)")
    settings.rate_limit = "10000/minute"
    with TestClient(main.app) as c:
        yield c

def convert(c, tool, fixture=None, options=None, url=None, assets=(), expect=200):
    """`fixture` is a file name in fixtures/ or a (name, bytes) pair."""
    files = []
    if isinstance(fixture, tuple):
        files = [("file", fixture)]
    elif fixture:
        files = [("file", (fixture, (FIX / fixture).read_bytes()))]
    files += [("assets", a) for a in assets]
    data = {"options": json.dumps(options or {})}
    if url:
        data["url"] = url
    r = c.post(f"/api/convert/{tool}", files=files or None, data=data)
    assert r.status_code == expect, (tool, r.status_code, r.text[:300])
    return r

def pdf_text(data: bytes) -> str:
    with fitz.open(stream=data, filetype="pdf") as d:
        return "".join(p.get_text() for p in d)

def docx_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read("word/document.xml").decode("utf-8")

def no_secrets(text: str):
    for s in SECRETS:
        assert s not in text, f"leaked {s!r}"

def test_capabilities_list_every_tool(client):
    assert sorted(client.get("/api/capabilities").json()["tools"]) == sorted(main.TOOLS)

# ---- Office / HTML → PDF (Gotenberg) --------------------------------------------------------

def test_word_to_pdf_keeps_arabic(client):
    r = convert(client, "word-to-pdf", "report.docx")
    assert r.content.startswith(b"%PDF")
    text = pdf_text(r.content)
    assert "تقرير الصيانة الشهري" in text and "مكيف" in text
    assert "report.pdf" in r.headers["content-disposition"]

def test_excel_to_pdf_calculates_formulas(client):
    text = pdf_text(convert(client, "excel-to-pdf", "sales.xlsx", {"one_page": True}).content)
    assert "فبراير" in text and "7350" in text

def test_powerpoint_to_pdf_one_page_per_slide(client):
    data = convert(client, "powerpoint-to-pdf", "slides.pptx").content
    with fitz.open(stream=data, filetype="pdf") as d:
        assert d.page_count == 2
    assert "تقرير الصيانة الشهري" in pdf_text(data)

def test_convert_anything_to_pdf(client):
    txt = "تقرير الصيانة الشهري" + chr(10) + "Office Lab" + chr(10)
    txt = txt.encode("utf-8")
    assert "تقرير الصيانة الشهري" in pdf_text(convert(client, "convert-to-pdf", ("notes.txt", txt)).content)
    assert "Office Lab" in pdf_text(convert(client, "convert-to-pdf", "slides.pptx").content)
    assert convert(client, "convert-to-pdf", ("x.exe", b"MZ"), expect=415).json() == {"code": "wrong-type"}


def test_pdf_to_excel_tables(client):
    from openpyxl import load_workbook

    table_pdf = convert(client, "excel-to-pdf", "sales.xlsx").content
    wb = load_workbook(io.BytesIO(convert(client, "pdf-to-excel", ("sales.pdf", table_pdf)).content))
    cells = [c.value for ws in wb for row in ws.iter_rows() for c in row]
    assert "فبراير" in cells and 2750 in cells
    assert "الإجمالي" in cells  # the lam-alef + hamza ligature comes back in the right order
    assert [c.value for c in wb.worksheets[0][1]] == ["الشهر", "المبلغ"]  # Arabic table: first column on the right
    assert wb.worksheets[0].sheet_view.rightToLeft
    assert convert(client, "pdf-to-excel", "text.pdf", expect=422).json() == {"code": "nothing"}


def test_pdf_to_excel_never_writes_formulas(tmp_path):
    """Text that looks like a formula in a PDF must land in Excel as plain text."""
    import importlib.util

    from openpyxl import Workbook, load_workbook

    spec = importlib.util.spec_from_file_location("pdf_to_xlsx", Path(main.__file__).parent / "scripts" / "pdf_to_xlsx.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    wb = Workbook()
    for i, text in enumerate(['=HYPERLINK("http://x","click")', "=1+1", "+cmd|' /C calc'!A0", "@SUM(1)"], start=1):
        mod.put(wb.active, i, 1, mod.value(text))
    wb.save(tmp_path / "f.xlsx")
    back = load_workbook(tmp_path / "f.xlsx").active
    assert [c.data_type for c in back["A"]] == ["s"] * 4
    assert back["A1"].value.startswith("=HYPERLINK")


def test_word_to_pdfa(client):
    assert b"pdfaid" in convert(client, "word-to-pdf", "report.docx", {"pdfa": True}).content

def test_html_to_pdf_with_assets(client):
    html = b'<!doctype html><link rel="stylesheet" href="style.css"><h1>Hello</h1>'
    css = b'h1::after{content:" ASSET-OK"}'
    r = convert(client, "html-to-pdf", ("index.html", html), assets=[("style.css", css)])
    assert "ASSET-OK" in pdf_text(r.content)
    assert "تقرير الصيانة الشهري" in pdf_text(convert(client, "html-to-pdf", "page.html").content)

def test_html_cannot_reach_internal_services_or_files(client):
    r = convert(client, "html-to-pdf", "hostile.html")
    no_secrets(pdf_text(r.content))

def test_html_cannot_list_other_jobs(client):
    """Gotenberg keeps each request to its own folder: no listing of /tmp or sibling jobs."""
    html = ('<!doctype html><h1>probe</h1><iframe src="../"></iframe><iframe src="../../"></iframe>'
            '<iframe src="file:///tmp/"></iframe><object data="../../"></object>').encode()
    text = pdf_text(convert(client, "html-to-pdf", ("index.html", html)).content)
    assert "probe" in text and "Index of" not in text


def test_url_to_pdf_and_image(client):
    assert "documentation examples" in pdf_text(convert(client, "url-to-pdf", url="https://example.com").content)
    png = convert(client, "url-to-image", url="https://example.com", options={"device": "phone"}).content
    assert png.startswith(b"\x89PNG")

@pytest.mark.parametrize("url", ["http://api:8000/api/capabilities", "http://127.0.0.1:3000/health",
                                 "http://localtest.me:8000/api/capabilities", "http://gotenberg:3000/health"])
def test_url_tools_refuse_internal_addresses(client, url):
    r = client.post("/api/convert/url-to-pdf", data={"url": url})
    assert r.status_code != 200
    if r.headers.get("content-type", "").startswith("application/pdf"):
        no_secrets(pdf_text(r.content))

def test_pdf_to_pdfa(client):
    assert b"pdfaid" in convert(client, "pdf-to-pdfa", "text.pdf", {"level": "PDF/A-2b"}).content

# ---- PDF (Ghostscript, qpdf, OCRmyPDF, pdf2docx) --------------------------------------------

def test_strong_compress_levels(client):
    original = (FIX / "photo.pdf").stat().st_size
    small = convert(client, "compress-pdf-strong", "photo.pdf", {"level": "screen"})
    assert len(small.content) < original * 0.5
    assert "x-office-lab-note" not in small.headers

def test_strong_compress_target_size(client):
    r = convert(client, "compress-pdf-strong", "photo.pdf", {"mode": "target", "target_mb": 0.15})
    assert len(r.content) <= 0.15 * 1024 * 1024 or r.headers.get("x-office-lab-note") == "target-missed"
    r = convert(client, "compress-pdf-strong", "photo.pdf", {"mode": "target", "target_mb": 0.0501})
    assert r.headers.get("x-office-lab-note") in ("target-missed", None)

def test_compress_never_returns_bigger(client):
    r = convert(client, "compress-pdf-strong", "text.pdf", {"level": "printer"})
    assert len(r.content) <= (FIX / "text.pdf").stat().st_size

def test_repair(client):
    data = convert(client, "repair-pdf", "damaged.pdf").content
    with fitz.open(stream=data, filetype="pdf") as d:
        assert d.page_count == 2 and not d.is_repaired

def test_password_protected_pdf(client):
    for tool in ("pdf-to-word", "compress-pdf-strong", "ocr-pdf-server"):
        assert convert(client, tool, "locked.pdf", expect=422).json() == {"code": "password"}

def test_pdf_to_word(client):
    xml = docx_xml(convert(client, "pdf-to-word", "text.pdf").content)
    assert "Office Lab test page 1" in xml and "lazy dog" in xml

def poppler_text(data: bytes) -> str:
    """PyMuPDF returns OCR'd Arabic in drawing order; poppler (like PDF viewers) gives reading order."""
    return subprocess.run(["pdftotext", "-", "-"], input=data, capture_output=True, check=True).stdout.decode("utf-8")

def test_ocr_pdf_arabic(client):
    text = poppler_text(convert(client, "ocr-pdf-server", "scan.pdf", {"lang": "ara"}).content)
    words = "تمت صيانة وحدات التكييف في المبنى الإداري".split()
    found = sum(w in text for w in words)
    assert found >= len(words) - 1, text

# ---- Ebooks (calibre) and documents (pandoc) ------------------------------------------------

@pytest.fixture(scope="module")
def epub(client):
    r = convert(client, "document-convert", "notes.md", {"to": "epub"})
    assert r.content.startswith(b"PK")
    return r.content

def test_ebook_chain(client, epub):
    azw3 = convert(client, "ebook-convert", ("notes.epub", epub), {"to": "azw3"}).content
    assert b"BOOKMOBI" in azw3[:100]
    back = convert(client, "mobi-to-epub", ("notes.azw3", azw3)).content
    assert back.startswith(b"PK")
    pdf = convert(client, "epub-to-pdf", ("notes.epub", epub)).content
    assert "تقرير الصيانة الشهري" in poppler_text(pdf)

def test_pdf_to_epub(client):
    assert convert(client, "pdf-to-epub", "text.pdf").content.startswith(b"PK")

def test_ebook_cannot_read_server_files(client):
    for to in ("txt", "docx", "fb2"):
        data = convert(client, "ebook-convert", "hostile.epub", {"to": to}).content
        text = data.decode("utf-8", "ignore") if to in ("txt", "fb2") else ""
        if to == "docx":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                text = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
        assert "after" in text  # the book itself converted
        no_secrets(text)

def test_markdown_to_word_is_rtl(client):
    xml = docx_xml(convert(client, "document-convert", "notes.md", {"to": "docx"}).content)
    assert "تقرير الصيانة الشهري" in xml and "w:bidi" in xml

def test_latex_word_both_ways(client):
    assert "Intro" in docx_xml(convert(client, "latex-to-word", "paper.tex").content)
    tex = convert(client, "word-to-latex", "report.docx").content.decode("utf-8")
    assert "\\begin{document}" in tex and "تقرير الصيانة الشهري" in " ".join(tex.split())

def test_pandoc_cannot_read_server_files(client):
    for tool, fixture, opts in (("latex-to-word", "hostile.tex", {}), ("document-convert", "hostile.md", {"to": "html"}),
                                ("document-convert", "hostile.md", {"to": "docx"})):
        r = client.post(f"/api/convert/{tool}", files={"file": (fixture, (FIX / fixture).read_bytes())},
                        data={"options": json.dumps(opts)})
        if r.status_code == 200:
            body = r.content
            if body.startswith(b"PK"):
                with zipfile.ZipFile(io.BytesIO(body)) as z:
                    body = b"".join(z.read(n) for n in z.namelist())
            no_secrets(body.decode("utf-8", "ignore"))

def test_same_format_is_refused(client, epub):
    assert convert(client, "ebook-convert", ("a.epub", epub), {"to": "epub"}, expect=422).json() == {"code": "bad-options"}
