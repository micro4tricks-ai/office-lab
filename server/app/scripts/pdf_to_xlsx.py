"""Tables in a PDF → an Excel workbook (one sheet per table). Run as its own process:

    pdf_to_xlsx.py in.pdf out.xlsx      exit 3 = no table found
"""

import re
import sys

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

NUMBER = re.compile(r"^-?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")
ARABIC = re.compile(r"[؀-ۿ]")
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")


def value(text: str | None):
    """Numbers become numbers (so sums work) — except ones with a leading zero (phones, codes)."""
    t = " ".join((text or "").split())
    plain = t.translate(DIGITS)
    if NUMBER.match(plain) and not (len(plain) > 1 and plain.startswith("0") and not plain.startswith("0.")):
        n = float(plain.replace(",", ""))
        return int(n) if n.is_integer() and "." not in plain else n
    return t


# Lam-alef ligatures with a hamza come out of many PDFs as "اإل" (alef, hamza-alef, lam) instead of "الإ".
LIGATURE = re.compile(r"ا([إأآ])ل")


def cell_text(words, bbox) -> str:
    """Words whose centre is in the cell, in reading order (right to left for Arabic lines)."""
    if bbox is None:
        return ""
    x0, y0, x1, y1 = bbox
    inside = [w for w in words if x0 <= (w[0] + w[2]) / 2 <= x1 and y0 <= (w[1] + w[3]) / 2 <= y1]
    lines: dict[tuple[int, int], list] = {}
    for w in inside:
        lines.setdefault((w[5], w[6]), []).append(w)
    out = []
    for key in sorted(lines, key=lambda k: min(w[1] for w in lines[k])):
        ws = lines[key]
        rtl = any(ARABIC.search(w[4]) for w in ws)
        ws.sort(key=lambda w: -w[0] if rtl else w[0])
        out.append(" ".join(w[4] for w in ws))
    return LIGATURE.sub(r"ال\1", " ".join(out)).strip()


def put(ws, row: int, col: int, v) -> None:
    """Writes a value; text always stays text. openpyxl would turn "=..." into a live formula,
    so a PDF could plant formulas (=HYPERLINK, DDE) in the workbook."""
    cell = ws.cell(row=row, column=col, value=v)
    if isinstance(v, str):
        cell.data_type = "s"


def main(src: str, dst: str) -> int:
    wb = Workbook()
    wb.remove(wb.active)
    with pymupdf.open(src) as doc:
        for pno, page in enumerate(doc, start=1):
            words = page.get_text("words")
            # Ruled tables first; tables without lines (e.g. printed from Excel) by text alignment.
            tables = page.find_tables().tables or page.find_tables(strategy="text").tables
            for tno, table in enumerate(tables, start=1):
                rows = [[cell_text(words, bbox) for bbox in row.cells] for row in table.rows]
                rows = [r for r in rows if any(r)]
                if not rows or len(rows) < 2 and len(rows[0]) < 2:
                    continue
                rtl = sum(bool(ARABIC.search(c)) for row in rows for c in row) > sum(bool(c) for row in rows for c in row) / 3
                ws = wb.create_sheet(f"p{pno}-t{tno}"[:31])
                for r, row in enumerate(rows, start=1):
                    # An Arabic table reads from the right: its first column is the rightmost one.
                    for c, cell in enumerate(reversed(row) if rtl else row, start=1):
                        put(ws, r, c, value(cell))
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for col in ws.columns:
                    width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                    ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
                    for c in col:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.sheet_view.rightToLeft = rtl
    if not wb.sheetnames:
        return 3
    wb.save(dst)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
